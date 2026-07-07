from collections import Counter
from collections import defaultdict
from datetime import datetime
from io import BytesIO
from datetime import date
from datetime import timedelta
import pandas as pd
from fastapi import APIRouter, Depends, File, Form, Request, UploadFile
from fastapi.responses import RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy.orm import Session
import re, calendar, io, math
from app.database import SessionLocal
from app.models.bom import BOM
from app.models.inventory import Inventory
from app.models.production_plan import ProductionPlan
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side
)
from openpyxl.utils import get_column_letter
from app.models.material_master import MaterialMaster
from app.models.material_note import MaterialNote
from calendar import monthcalendar
from app.utils.logger import save_log
router = APIRouter()
templates = Jinja2Templates(directory="app/templates")

MRP_LOG_PRODUCT = "MRP"


def current_user(request: Request):
    return request.session.get("user") or "system"

BOM_COLUMNS = {
    "product_name": ["Product", "product_name", "제품명"],
    "component_code": ["Item Code", "component_code", "품목코드"],
    "component_name": ["Item Name", "component_name", "품목명"],
    "qty": ["Qty", "qty", "소요량"],
}


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def find_column(df, candidates):
    columns = {
        str(column).strip().lower(): column
        for column in df.columns
    }

    for candidate in candidates:
        column = columns.get(candidate.lower())
        if column is not None:
            return column

    return None


def normalize_bom_dataframe(df):
    mapped = {
        key: find_column(df, candidates)
        for key, candidates in BOM_COLUMNS.items()
    }

    missing = [
        BOM_COLUMNS[key][0]
        for key, value in mapped.items()
        if value is None
    ]

    if missing:
        raise ValueError(
            "Missing BOM columns: " + ", ".join(missing)
        )

    rows = df.rename(
        columns={
            mapped["product_name"]: "product_name",
            mapped["component_code"]: "component_code",
            mapped["component_name"]: "component_name",
            mapped["qty"]: "qty",
        }
    )

    rows = rows[
        ["product_name", "component_code", "component_name", "qty"]
    ].copy()
    rows = rows.dropna(
        subset=["product_name", "component_code", "qty"]
    )

    rows["product_name"] = rows["product_name"].astype(str).str.strip()
    rows["component_code"] = rows["component_code"].astype(str).str.strip()
    rows["component_name"] = (
        rows["component_name"].fillna("").astype(str).str.strip()
    )
    rows["qty"] = pd.to_numeric(rows["qty"], errors="coerce").fillna(0)

    return rows[rows["qty"] > 0]


def get_bom_rows(db):
    return (
        db.query(BOM)
        .order_by(BOM.product_name, BOM.component_code)
        .all()
    )

def get_week_label(dt):
    cal = calendar.monthcalendar(dt.year, dt.month)

    for idx, week in enumerate(cal, start=1):
        if dt.day in week:
            return f"{dt.month}월 {idx}주"

    return f"{dt.month}월 1주"

def calculate_mrp(
    db,
    year=None,
    month=None,
    week=None,
    all_plans=None,
    bom_rows=None,
    inventory_rows=None,
    material_rows=None,
    note_rows=None,
):
    # all_plans/bom_rows/inventory_rows/material_rows/note_rows는 호출자가
    # 이미 조회해둔 데이터가 있으면 그대로 재사용하기 위한 선택 인자다 -
    # 넘기지 않으면(None) 예전과 동일하게 이 함수가 직접 조회한다. 계산
    # 로직은 그대로고, 같은 테이블을 여러 번 중복 조회하지 않기 위함.
    if all_plans is None:
        all_plans = (
            db.query(ProductionPlan)
            .order_by(
                ProductionPlan.plan_date,
                ProductionPlan.product_name
            )
            .all()
        )

    def matches_period(plan):
        if year and plan.plan_date.year != year:
            return False

        if month and plan.plan_date.month != month:
            return False

        if week and plan.plan_date.isocalendar().week != week:
            return False

        return True

    filtered_plans = [
        plan for plan in all_plans
        if matches_period(plan)
    ]

    plans = (
        filtered_plans
        if (year or month or week)
        else all_plans
    )

    products = {
        p.product_name
        for p in all_plans
    }

    if bom_rows is None:
        bom_rows = (
            db.query(BOM)
            .filter(
                BOM.product_name.in_(products)
            )
            .all()
            if products
            else []
        )
    else:
        bom_rows = [
            bom for bom in bom_rows
            if bom.product_name in products
        ]

    bom_by_product = defaultdict(list)

    for bom in bom_rows:
        bom_by_product[
            bom.product_name
        ].append(bom)

    required_by_item = {}

    for plan in all_plans:

        bom_list = bom_by_product.get(
            plan.product_name,
            []
        )

        if not bom_list:
            continue

        for bom in bom_list:

            required_qty = (
                (plan.plan_qty or 0)
                *
                (bom.qty or 0)
            )

            item = required_by_item.setdefault(
                bom.component_code,
                {
                    "item_code": bom.component_code,
                    "item_name": bom.component_name,
                    "required_qty": 0,
                    "products": defaultdict(float),
                    "details": []
                }
            )

            item["required_qty"] += required_qty

            item["products"][
                plan.product_name
            ] += required_qty

            item["details"].append(
                {
                    "date": plan.plan_date,
                    "year": plan.plan_date.year,
                    "month": plan.plan_date.month,
                    "week": get_week_label(plan.plan_date),
                    "in_period": matches_period(plan),
                    "product_name": plan.product_name,
                    "plan_qty": plan.plan_qty or 0,
                    "bom_qty": bom.qty or 0,
                    "required_qty": required_qty,
                }
            )

    for item in required_by_item.values():

        item["details"].sort(
            key=lambda x: x["date"]
        )

    if inventory_rows is None:
        inventory_rows = db.query(Inventory).all()

    if material_rows is None:
        material_rows = db.query(MaterialMaster).all()

    if note_rows is None:
        note_rows = db.query(MaterialNote).all()

    material_map = {
        row.item_code: row
        for row in material_rows
    }

    note_map = {
        row.item_code: row
        for row in note_rows
    }

    inventory_by_code = defaultdict(list)

    for inv in inventory_rows:

        inventory_by_code[
            inv.item_code
        ].append(inv)

    results = []

    for item_code, row in required_by_item.items():

        inventory_list = inventory_by_code.get(
            item_code,
            []
        )

        material = material_map.get(
            item_code
        )

        note_row = note_map.get(
            item_code
        )

        note = (
            note_row.note
            if note_row
            else ""
        )

        has_note = bool(
            note.strip()
        )

        supplier = (
            material.supplier
            if material
            else ""
        )

        lead_time_week = (
            material.lead_time_week
            if material
            else 0
        )

        moq = (
            material.moq
            if material
            else 0
        )

        inventory_detail = {
            "창고재고": 0,
            "제공재고": 0,
            "외주재고": 0,
        }

        stock_qty = 0

        for inv in inventory_list:

            qty = inv.qty or 0

            if inv.warehouse_type == "창고재고":

                if inv.grade == "A":
                    inventory_detail["창고재고"] += qty
                    stock_qty += qty

            elif inv.warehouse_type == "제공재고":

                inventory_detail["제공재고"] += qty
                stock_qty += qty

            elif inv.warehouse_type == "외주재고":

                inventory_detail["외주재고"] += qty
                stock_qty += qty

        period_details = [
            detail for detail in row["details"]
            if detail.get("in_period")
        ]

        if (year or month or week) and not period_details:
            continue

        display_details = (
            period_details
            if (year or month or week)
            else row["details"]
        )

        period_start = (
            min(
                detail["date"]
                for detail in display_details
            )
            if display_details
            else None
        )

        consumed_before_period = (
            sum(
                detail["required_qty"]
                for detail in row["details"]
                if period_start
                and detail["date"] < period_start
            )
            if (year or month or week)
            else 0
        )

        available_stock_qty = max(
            stock_qty - consumed_before_period,
            0
        )

        required_qty = sum(
            detail["required_qty"]
            for detail in display_details
        )

        shortage_qty = max(
            required_qty - available_stock_qty,
            0
        )

        if shortage_qty > 0 and moq > 0:

            recommended_order_qty = (
                math.ceil(
                    shortage_qty / moq
                ) * moq
            )

        else:

            recommended_order_qty = shortage_qty

        coverage_rate = (
            round(
                min(
                    available_stock_qty / required_qty * 100,
                    100
                ),
                1
            )
            if required_qty > 0
            else 100
        )

        product_names = sorted(
            {
                detail["product_name"]
                for detail in display_details
            }
        )

        remain_stock = available_stock_qty

        need_date = None

        weekly_required = defaultdict(float)

        for detail in display_details:

            weekly_required[
                detail["date"]
            ] += detail["required_qty"]

        for target_date, required in sorted(
            weekly_required.items()
        ):

            if remain_stock >= required:

                remain_stock -= required

            else:

                need_date = target_date

                break

        if need_date is None and weekly_required:

            need_date = max(
                weekly_required.keys()
            )

        order_date = None

        if need_date:

            order_date = (
                need_date
                - timedelta(
                    weeks=lead_time_week
                )
            )

        if shortage_qty > 0:

            priority = (
                "Critical"
                if available_stock_qty == 0
                else "Short"
            )

        else:

            priority = "Ready"

        results.append(
            {
                "item_code": item_code,
                "item_name": row["item_name"],
                "product_count": len(product_names),
                "product_names": product_names,
                "product_summary": (
                    product_names[0]
                    if len(product_names) == 1
                    else f"{len(product_names)} products"
                ),
                "warehouse_stock":
                    inventory_detail["창고재고"],
                "supplier_stock":
                    inventory_detail["제공재고"],
                "subcontract_stock":
                    inventory_detail["외주재고"],
                "details": sorted(
                    display_details,
                    key=lambda x: x["date"]
                ),
                "required_qty": required_qty,
                "stock_qty": available_stock_qty,
                "physical_stock_qty": stock_qty,
                "consumed_before_period": consumed_before_period,
                "shortage_qty": shortage_qty,
                "coverage_rate": coverage_rate,
                "priority": priority,
                "status": (
                    "SHORT"
                    if shortage_qty > 0
                    else "OK"
                ),
                "supplier": supplier,
                "lead_time_week": lead_time_week,
                "moq": moq,
                "need_date": need_date,
                "order_date": order_date,
                "recommended_order_qty":
                    recommended_order_qty,
                "note": note,
                "has_note": has_note
            }
        )
    
    selected_year = year or datetime.now().year
    selected_month = month or datetime.now().month

    week_summary = defaultdict(
        lambda: {
            "plan_qty": 0,
            "required_qty": 0,
            "stock_qty": 0,
            "coverage": 100,
        }
    )

    for plan in plans:

        label = get_week_label(
            plan.plan_date
        )

        week_summary[label]

        week_summary[label]["plan_qty"] += (
            plan.plan_qty or 0
        )

    for row in results:

        remain_stock = row["stock_qty"]

        weekly_required = defaultdict(float)

        for detail in row["details"]:

            required = detail["required_qty"]

            if remain_stock >= required:
                supplied = required
                remain_stock -= required
            else:
                supplied = remain_stock
                remain_stock = 0

            label = get_week_label(detail["date"])

            if label in week_summary:

                week_summary[label]["required_qty"] += required

                week_summary[label]["stock_qty"] += min(
                    supplied,
                    required
                )

    for label, data in week_summary.items():

        if data["required_qty"] == 0:
            data["coverage"] = 0

        else:
            data["coverage"] = round(
                min(
                    data["stock_qty"]
                    /
                    data["required_qty"]
                    * 100,
                    100
                ),
                1
            )

    results.sort(
        key=lambda row: (
            row["priority"] != "Critical",
            row["priority"] != "Short",
            row["need_date"] or date.max,
            row["item_code"]
        )
    )

    shortage_count = sum(
        1
        for row in results
        if row["shortage_qty"] > 0
    )

    return (
        plans,
        results,
        shortage_count,
        week_summary
    )


def filter_mrp_rows(rows, q="", shortage_only=False):
    keyword = (q or "").strip().lower()

    if keyword:
        rows = [
            row for row in rows
            if keyword in (row["item_code"] or "").lower()
            or keyword in (row["item_name"] or "").lower()
            or any(
                keyword in product.lower()
                for product in row["product_names"]
            )
        ]

    if shortage_only:
        rows = [
            row for row in rows
            if row["shortage_qty"] > 0
        ]

    return rows

def build_week_dashboard_summary(
    db,
    plans,
    year=None,
    month=None,
    week=None,
    inventory_rows=None,
    bom_rows=None,
):

    if inventory_rows is None:
        inventory_rows = db.query(Inventory).all()

    if bom_rows is None:
        bom_rows = db.query(BOM).all()

    remain_inventory = defaultdict(float)

    for inv in inventory_rows:

        qty = inv.qty or 0

        if inv.warehouse_type == "창고재고":
            if inv.grade == "A":
                remain_inventory[inv.item_code] += qty

        elif inv.warehouse_type in ["제공재고", "외주재고"]:
            remain_inventory[inv.item_code] += qty

    bom_map = defaultdict(list)

    for bom in bom_rows:
        bom_map[bom.product_name].append({
            "item_code": bom.component_code,
            "qty": bom.qty or 0
        })

    plans = sorted(
        plans,
        key=lambda p: p.plan_date
    )

    def matches_period(plan):
        if year and plan.plan_date.year != year:
            return False

        if month and plan.plan_date.month != month:
            return False

        if week and plan.plan_date.isocalendar().week != week:
            return False

        return True

    week_summary = defaultdict(
        lambda: {
            "plan_qty": 0,
            "possible_qty": 0,
            "coverage": 0
        }
    )

    for plan in plans:

        label = get_week_label(plan.plan_date)

        in_period = matches_period(plan)

        if in_period:

            week_summary[label]["plan_qty"] += (
                plan.plan_qty or 0
            )

        bom_list = bom_map.get(
            plan.product_name,
            []
        )

        if not bom_list:
            continue

        possible_qty = math.inf

        for bom in bom_list:

            if bom["qty"] <= 0:
                continue

            stock = remain_inventory.get(
                bom["item_code"],
                0
            )

            possible_qty = min(
                possible_qty,
                stock / bom["qty"]
            )

        if possible_qty == math.inf:
            possible_qty = 0

        possible_qty = max(
            0,
            math.floor(possible_qty)
        )

        produced_qty = int(
            min(
                possible_qty,
                plan.plan_qty or 0
            )
        )

        if in_period:

            week_summary[label]["possible_qty"] += (
                produced_qty
            )

        for bom in bom_list:

            remain_inventory[
                bom["item_code"]
            ] -= (
                produced_qty
                *
                bom["qty"]
            )

    for label, data in week_summary.items():

        if data["plan_qty"] > 0:

            data["coverage"] = round(
                min(
                    data["possible_qty"]
                    /
                    data["plan_qty"]
                    * 100,
                    100
                ),
                1
            )

        else:

            data["coverage"] = 0

    return week_summary

def build_mrp_summary(plans, rows, filtered_rows, inventory_count):
    shortage_rows = [
        row for row in rows
        if row["shortage_qty"] > 0
    ]

    total_required = sum(
        row["required_qty"]
        for row in rows
    )

    total_stock = sum(
        min(
            row["stock_qty"],
            row["required_qty"]
        )
        for row in rows
    )

    total_shortage = (
        total_required
        - total_stock
    )

    return {
        "plan_count": sum(
            plan.plan_qty or 0
            for plan in plans
        ),
        "required_item_count": len(rows),
        "filtered_item_count": len(filtered_rows),
        "shortage_count": len(shortage_rows),
        "ready_count": len(rows) - len(shortage_rows),
        "inventory_count": inventory_count,
        "total_required": total_required,
        "total_stock": total_stock,
        "total_shortage": total_shortage,
        "shortage_rate": (
            len(shortage_rows) / len(rows) * 100
            if rows
            else 0
        ),
        "coverage_rate": (
            round(
                min(total_stock / total_required * 100, 100),
                1
            )
            if total_required > 0
            else 0
        ),
        "calculated_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }

def build_mrp_dashboard_context(
    db,
    plans,
    rows,
    shortage_count,
    week_summary,
    selected_year=None,
    selected_month=None,
    inventory_rows=None,
    bom_rows=None,
    material_rows=None,
    note_rows=None,
):

    if inventory_rows is None:
        inventory_rows = db.query(Inventory).all()

    if bom_rows is None:
        bom_rows = db.query(BOM).all()

    if material_rows is None:
        material_rows = db.query(MaterialMaster).all()

    note_count = (
        db.query(MaterialNote).count()
        if note_rows is None
        else len(note_rows)
    )

    inventory_count = len(inventory_rows)
    summary = build_mrp_summary(
        plans,
        rows,
        rows,
        inventory_count,
    )

    shortage_rows = [
        row for row in rows
        if row["shortage_qty"] > 0
    ]
    ready_rows = [
        row for row in rows
        if row["shortage_qty"] <= 0
    ]
    critical_rows = [
        row for row in shortage_rows
        if row["priority"] == "Critical"
    ]

    total_required = summary["total_required"]
    total_stock = summary["total_stock"]
    total_shortage = summary["total_shortage"]
    total_recommended_order = sum(
        row["recommended_order_qty"]
        for row in shortage_rows
    )
    shortage_value_rate = (
        total_shortage / total_required * 100
        if total_required
        else 0
    )

    material_codes = {
        row.item_code
        for row in material_rows
    }
    required_codes = {
        row["item_code"]
        for row in rows
    }
    missing_master_count = len(
        required_codes - material_codes
    )

    supplier_rows = {}
    for row in shortage_rows:
        supplier = row["supplier"] or "미지정"
        if supplier not in supplier_rows:
            supplier_rows[supplier] = {
                "name": supplier,
                "item_count": 0,
                "shortage_qty": 0,
                "order_qty": 0,
            }
        supplier_rows[supplier]["item_count"] += 1
        supplier_rows[supplier]["shortage_qty"] += row["shortage_qty"]
        supplier_rows[supplier]["order_qty"] += row["recommended_order_qty"]

    supplier_summary = sorted(
        supplier_rows.values(),
        key=lambda row: row["shortage_qty"],
        reverse=True,
    )[:8]

    product_rows = {}
    for row in rows:
        product_names = row["product_names"] or ["미지정"]
        for product in product_names:
            if product not in product_rows:
                product_rows[product] = {
                    "product": product,
                    "required_qty": 0,
                    "shortage_items": 0,
                    "shortage_qty": 0,
                }
            product_rows[product]["required_qty"] += (
                row["required_qty"] / len(product_names)
                if product_names
                else row["required_qty"]
            )
            if row["shortage_qty"] > 0:
                product_rows[product]["shortage_items"] += 1
                product_rows[product]["shortage_qty"] += (
                    row["shortage_qty"] / len(product_names)
                    if product_names
                    else row["shortage_qty"]
                )

    product_summary = sorted(
        product_rows.values(),
        key=lambda row: (
            row["shortage_items"],
            row["shortage_qty"],
        ),
        reverse=True,
    )[:8]

    plan_by_month = {}

    for plan in plans:

        if selected_year and plan.plan_date.year != selected_year:
            continue

        if selected_month and plan.plan_date.month != selected_month:
            continue

        label = plan.plan_date.strftime("%Y-%m")

        if label not in plan_by_month:
            plan_by_month[label] = {
                "qty": 0,
                "count": 0,
            }

        plan_by_month[label]["qty"] += plan.plan_qty or 0
        plan_by_month[label]["count"] += 1

    plan_labels = sorted(plan_by_month.keys())
    plan_qtys = [
        plan_by_month[label]["qty"]
        for label in plan_labels
    ]
    plan_counts = [
        plan_by_month[label]["count"]
        for label in plan_labels
    ]

    inventory_by_warehouse = Counter()
    for row in inventory_rows:
        inventory_by_warehouse[row.warehouse_type or "미지정"] += (
            row.qty or 0
        )

    bom_by_product = Counter()
    bom_qty_by_product = Counter()
    for row in bom_rows:
        bom_by_product[row.product_name or "미지정"] += 1
        bom_qty_by_product[row.product_name or "미지정"] += row.qty or 0

    bom_top = [
        {
            "product": product,
            "component_count": count,
            "total_qty": bom_qty_by_product[product],
        }
        for product, count in bom_by_product.most_common(8)
    ]

    week_labels = []
    week_plan_qty = []
    week_possible_qty = []
    week_coverage = []

    if selected_year and selected_month:

        cal = calendar.monthcalendar(
            selected_year,
            selected_month
        )

        for i in range(len(cal)):

            label = f"{selected_month}월 {i+1}주"

            data = week_summary.get(
                label,
                {
                    "plan_qty": 0,
                    "possible_qty": 0,
                    "coverage": 0
                }
            )

            week_labels.append(label)
            week_plan_qty.append(data["plan_qty"])
            week_possible_qty.append(data["possible_qty"])
            week_coverage.append(round(data["coverage"], 1))

    else:

        for label in sorted(
            week_summary.keys(),
            key=lambda x: (
                int(x.split("월")[0]),
                int(x.split(" ")[1].replace("주", ""))
            )
        ):

            data = week_summary[label]

            week_labels.append(label)
            week_plan_qty.append(data["plan_qty"])
            week_possible_qty.append(data["possible_qty"])
            week_coverage.append(round(data["coverage"], 1))

    urgent_orders = sorted(
        shortage_rows,
        key=lambda row: (
            row["order_date"] is None,
            row["order_date"] or date.max,
            -row["shortage_qty"],
        ),
    )[:8]

    upcoming_plans = sorted(
        plans,
        key=lambda row: row.plan_date,
    )[:8]

    top_shortages = sorted(
        shortage_rows,
        key=lambda row: row["shortage_qty"],
        reverse=True,
    )[:10]

    supplier_count = len(
        {
            row.supplier
            for row in material_rows
            if row.supplier
        }
    )

    return {
        "summary": summary,
        "bom_count": len(bom_rows),
        "bom_product_count": len(bom_by_product),
        "plan_count": len(plans),
        "inventory_count": inventory_count,
        "material_count": len(material_rows),
        "supplier_count": supplier_count,
        "note_count": note_count,
        "shortage_count": shortage_count,
        "critical_count": len(critical_rows),
        "ready_count": len(ready_rows),
        "missing_master_count": missing_master_count,
        "total_required": total_required,
        "total_stock": total_stock,
        "total_shortage": total_shortage,
        "total_recommended_order": total_recommended_order,
        "shortage_value_rate": shortage_value_rate,
        "top_shortages": top_shortages,
        "urgent_orders": urgent_orders,
        "supplier_summary": supplier_summary,
        "product_summary": product_summary,
        "bom_top": bom_top,
        "upcoming_plans": upcoming_plans,
        "now": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "risk_labels": ["Critical", "Short", "Ready"],
        "risk_values": [
            len(critical_rows),
            len(shortage_rows) - len(critical_rows),
            len(ready_rows),
        ],
        "shortage_labels": [
            row["item_code"]
            for row in top_shortages[:8]
        ],
        "shortage_values": [
            round(row["shortage_qty"], 2)
            for row in top_shortages[:8]
        ],
        "supplier_labels": [
            row["name"]
            for row in supplier_summary
        ],
        "supplier_values": [
            round(row["shortage_qty"], 2)
            for row in supplier_summary
        ],
        "product_labels": [
            row["product"]
            for row in product_summary
        ],
        "product_values": [
            row["shortage_items"]
            for row in product_summary
        ],
        "plan_labels": plan_labels,
        "plan_qtys": plan_qtys,
        "week_possible_qty": week_possible_qty,
        "plan_counts": plan_counts,
        "inventory_labels": list(inventory_by_warehouse.keys()),
        "inventory_values": list(inventory_by_warehouse.values()),
        "week_labels": week_labels,
        "week_plan_qty": week_plan_qty,
        "week_coverage": week_coverage,
        "bom_labels": [
            row["product"]
            for row in bom_top
        ],
        "bom_values": [
            row["component_count"]
            for row in bom_top
        ],
    }

def parse_moq(value):

    if pd.isna(value):
        return 0

    value = str(value).strip().upper()

    if value in ["", "-"]:
        return 0

    if value.endswith("K"):

        return int(
            float(
                value.replace("K", "")
            ) * 1000
        )

    return int(
        float(value)
    )

@router.get("/mrp")
def mrp_dashboard(
    request: Request,
    db: Session = Depends(get_db),
):
    year = request.query_params.get("year")
    month = request.query_params.get("month")

    year = int(year) if year else None
    month = int(month) if month else None

    # 이 페이지는 calculate_mrp -> build_week_dashboard_summary ->
    # build_mrp_dashboard_context를 연달아 호출하는데, 셋 다 각자
    # ProductionPlan/Inventory/BOM/MaterialMaster/MaterialNote를 따로
    # 또 조회해서 같은 테이블을 최대 3번까지 중복으로 읽고 있었다.
    # 한 번씩만 가져와 세 함수에 그대로 넘겨 재사용한다 - 계산 로직과
    # 결과는 동일하고 중복 조회만 없앤다.
    all_plans = (
        db.query(ProductionPlan)
        .order_by(
            ProductionPlan.plan_date,
            ProductionPlan.product_name
        )
        .all()
    )
    inventory_rows = db.query(Inventory).all()
    bom_rows = db.query(BOM).all()
    material_rows = db.query(MaterialMaster).all()
    note_rows = db.query(MaterialNote).all()

    plans, rows, shortage_count, _ = calculate_mrp(
        db,
        year=year,
        month=month,
        all_plans=all_plans,
        bom_rows=bom_rows,
        inventory_rows=inventory_rows,
        material_rows=material_rows,
        note_rows=note_rows,
    )
    plans_for_filter = all_plans

    dashboard_week_summary = build_week_dashboard_summary(
        db,
        plans_for_filter,
        year=year,
        month=month,
        inventory_rows=inventory_rows,
        bom_rows=bom_rows,
    )
    context = build_mrp_dashboard_context(
        db,
        plans,
        rows,
        shortage_count,
        dashboard_week_summary,
        selected_year=year,
        selected_month=month,
        inventory_rows=inventory_rows,
        bom_rows=bom_rows,
        material_rows=material_rows,
        note_rows=note_rows,
    )

    years = sorted(
        {
            p.plan_date.year
            for p in plans_for_filter
        },
        reverse=True
    )

    if year:
        months = sorted(
            {
                p.plan_date.month
                for p in plans_for_filter
                if p.plan_date.year == year
            }
        )
    else:
        months = sorted(
            {
                p.plan_date.month
                for p in plans_for_filter
            }
        )

    context["selected_year"] = year
    context["selected_month"] = month

    context["years"] = years
    context["months"] = months

    return templates.TemplateResponse(
        request=request,
        name="bom/dashboard.html",
        context=context,
    )


@router.get("/mrp/bom")
@router.get("/mrp/bom/upload")
@router.get("/mrp/bom/list")
def bom_page(
    request: Request,
    db: Session = Depends(get_db),
):

    rows = get_bom_rows(db)

    grouped_rows = {}

    for row in rows:

        product_name = row.product_name

        if product_name not in grouped_rows:

            grouped_rows[product_name] = {
                "rows": [],
                "count": 0
            }

        grouped_rows[product_name]["rows"].append(
            row
        )

        grouped_rows[product_name]["count"] += 1

    return templates.TemplateResponse(
        request=request,
        name="bom/bom.html",
        context={
            "rows": rows,
            "grouped_rows": grouped_rows,
            "row_count": len(rows),
        },
    )


@router.post("/mrp/bom")
@router.post("/mrp/bom/upload")
async def upload_bom(
    request: Request,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    try:

        df = pd.read_excel(file.file)

        rows = normalize_bom_dataframe(df)

    except ValueError as exc:

        current_rows = get_bom_rows(db)

        grouped_rows = {}

        for row in current_rows:

            product_name = row.product_name

            if product_name not in grouped_rows:

                grouped_rows[product_name] = {
                    "rows": [],
                    "count": 0
                }

            grouped_rows[product_name]["rows"].append(
                row
            )

            grouped_rows[product_name]["count"] += 1

        return templates.TemplateResponse(
            request=request,
            name="bom/bom.html",
            context={
                "rows": current_rows,
                "grouped_rows": grouped_rows,
                "row_count": len(current_rows),
                "error": str(exc),
            },
            status_code=400,
        )

    db.query(BOM).delete()

    for row in rows.to_dict(
        orient="records"
    ):

        db.add(
            BOM(
                product_name=row["product_name"],
                component_code=row["component_code"],
                component_name=row["component_name"],
                qty=float(row["qty"]),
            )
        )

    db.commit()

    save_log(
        user=current_user(request),
        product=MRP_LOG_PRODUCT,
        action="MRP_BOM_UPLOAD",
        detail=f"BOM 업로드: {len(rows)}건"
    )

    current_rows = get_bom_rows(db)

    grouped_rows = {}

    for row in current_rows:

        product_name = row.product_name

        if product_name not in grouped_rows:

            grouped_rows[product_name] = {
                "rows": [],
                "count": 0
            }

        grouped_rows[product_name]["rows"].append(
            row
        )

        grouped_rows[product_name]["count"] += 1

    return templates.TemplateResponse(
        request=request,
        name="bom/bom.html",
        context={
            "rows": current_rows,
            "grouped_rows": grouped_rows,
            "row_count": len(current_rows),
            "message": f"Uploaded {len(rows)} BOM rows.",
        },
    )

@router.get("/mrp/bom/download")
def download_bom(
    request: Request,
    db: Session = Depends(get_db)
):

    rows = (
        db.query(BOM)
        .order_by(
            BOM.product_name,
            BOM.component_code
        )
        .all()
    )

    data = []

    for row in rows:

        data.append({

            "제품명": row.product_name,
            "품목코드": row.component_code,
            "품목명": row.component_name,
            "소요량": row.qty
        })

    df = pd.DataFrame(data)

    output = io.BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl"
    ) as writer:

        df.to_excel(
            writer,
            index=False,
            sheet_name="BOM"
        )

    output.seek(0)

    save_log(
        user=current_user(request),
        product=MRP_LOG_PRODUCT,
        action="MRP_BOM_DOWNLOAD",
        detail=f"BOM 다운로드: {len(rows)}건"
    )

    return StreamingResponse(

        output,

        media_type=
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",

        headers={
            "Content-Disposition":
            "attachment; filename=bom.xlsx"
        }

    )

@router.get("/mrp/production-plan")
def production_plan_page(
    request: Request,
    db: Session = Depends(get_db)
):
    keyword = (
        request.query_params.get(
            "keyword",
            ""
        )
        .strip()
    )

    start_date = request.query_params.get(
        "start_date"
    )

    end_date = request.query_params.get(
        "end_date"
    )

    products = (
        db.query(BOM.product_name)
        .distinct()
        .order_by(BOM.product_name)
        .all()
    )

    plans_query = db.query(
        ProductionPlan
    )

    if keyword:

        plans_query = plans_query.filter(
            ProductionPlan.product_name.contains(
                keyword
            )
        )

    if start_date:

        plans_query = plans_query.filter(
            ProductionPlan.plan_date >=
            datetime.strptime(
                start_date,
                "%Y-%m-%d"
            ).date()
        )

    if end_date:

        plans_query = plans_query.filter(
            ProductionPlan.plan_date <=
            datetime.strptime(
                end_date,
                "%Y-%m-%d"
            ).date()
        )

    plans = plans_query.all()

    is_search_mode = (
        bool(keyword)
        or bool(start_date)
        or bool(end_date)
    )

    if keyword or start_date or end_date:

        product_list = sorted(
            {
                p.product_name
                for p in plans
            }
        )

    else:

        product_list = [
            p[0]
            for p in products
        ]

    today = date.today()

    selected_year = int(
        request.query_params.get(
            "year",
            today.year
        )
    )

    selected_month = int(
        request.query_params.get(
            "month",
            today.month
        )
    )

    current_year = today.year

    year_columns = list(
        range(
            2026,
            current_year + 2
        )
    )

    month_columns = list(
        range(
            1,
            13
        )
    )

    week_columns = []

    cal = calendar.monthcalendar(
        selected_year,
        selected_month
    )

    for week in cal:

        first_day = next(
            (
                d
                for d in week
                if d != 0
            ),
            None
        )

        if first_day:

            iso = date(
                selected_year,
                selected_month,
                first_day
            ).isocalendar()

            item = (
                iso.year,
                iso.week
            )

            if item not in week_columns:

                week_columns.append(
                    item
                )
        
    if not week_columns:

        week_columns = [
            (
                today.isocalendar().year,
                today.isocalendar().week
            )
        ]

    selected_week = request.query_params.get(
        "week"
    )

    selected_iso_year = request.query_params.get(
        "iso_year"
    )

    if selected_week and selected_iso_year:

        selected_week = int(
            selected_week
        )

        selected_iso_year = int(
            selected_iso_year
        )

    else:

        selected_iso_year = (
            week_columns[0][0]
        )

        selected_week = (
            week_columns[0][1]
        )

    if (
        selected_iso_year,
        selected_week
    ) not in week_columns:

        selected_iso_year = (
            week_columns[0][0]
        )

        selected_week = (
            week_columns[0][1]
        )

    try:

        week_start = date.fromisocalendar(
            selected_iso_year,
            selected_week,
            1
        )

    except ValueError:

        selected_week = (
            week_columns[0]
            if week_columns
            else today.isocalendar()[1]
        )

        week_start = date.fromisocalendar(
            selected_year,
            selected_week,
            1
        )

    week_end = (
        week_start +
        timedelta(days=6)
    )

    if is_search_mode and plans:

        if start_date:

            display_start = datetime.strptime(
                start_date,
                "%Y-%m-%d"
            ).date()

        else:

            display_start = min(
                p.plan_date
                for p in plans
            )

        if end_date:

            display_end = datetime.strptime(
                end_date,
                "%Y-%m-%d"
            ).date()

        else:

            display_end = max(
                p.plan_date
                for p in plans
            )

        date_columns = [

            display_start +
            timedelta(days=i)

            for i in range(
                (
                    display_end -
                    display_start
                ).days + 1
            )

        ]

    else:

        display_start = week_start

        display_end = week_end

        date_columns = [

            week_start +
            timedelta(days=i)

            for i in range(7)

        ]

    if plans:

        df = pd.DataFrame([

            {
                "product_name": p.product_name,
                "plan_date": p.plan_date,
                "plan_qty": p.plan_qty
            }

            for p in plans

        ])

        pivot = df.pivot_table(
            index="product_name",
            columns="plan_date",
            values="plan_qty",
            aggfunc="sum",
            fill_value=0
        )

    else:

        pivot = pd.DataFrame()

    if not plans:

        return templates.TemplateResponse(
            request=request,
            name="bom/production_plan.html",
            context={
                "grouped_rows": {},
                "date_columns": date_columns,
                "year_columns": year_columns,
                "month_columns": month_columns,
                "week_columns": week_columns,
                "selected_year": selected_year,
                "selected_month": selected_month,
                "selected_week": selected_week,
                "selected_iso_year": selected_iso_year,
                "display_start": display_start,
                "display_end": display_end,
                "is_search_mode": is_search_mode,
                "keyword": keyword,
                "start_date": start_date,
                "end_date": end_date,
                "row_count": 0,
                "today_week": (
                    today.isocalendar()[1]
                )
            }
        )

    pivot = pivot.reindex(
        index=product_list,
        fill_value=0
    )

    pivot = pivot.reindex(
        columns=date_columns,
        fill_value=0
    )

    pivot = pivot.reset_index()

    if "product_name" not in pivot.columns:

        pivot = pivot.rename(
            columns={
                "index": "product_name"
            }
        )

    rows = pivot.to_dict(
        orient="records"
    )

    grouped_rows = {}

    for row in rows:

        product_name = row[
            "product_name"
        ]

        match = re.match(
            r"(.+?)_([0-9]+호)$",
            product_name
        )

        if match:

            group_name = (
                match.group(1)
            )

            size_name = (
                match.group(2)
            )

        else:

            group_name = product_name

            size_name = product_name

        row["display_name"] = (
            size_name
        )

        if group_name not in grouped_rows:

            grouped_rows[group_name] = {
                "rows": [],
                "total_qty": 0
            }

        weekly_total = 0

        for d in date_columns:

            weekly_total += (
                row.get(d, 0)
                or 0
            )

        row["weekly_total"] = (
            weekly_total
        )

        grouped_rows[group_name][
            "rows"
        ].append(
            row
        )

        grouped_rows[group_name][
            "total_qty"
        ] += weekly_total

    for group_name in grouped_rows:

        grouped_rows[group_name][
            "rows"
        ].sort(

            key=lambda x:

            int(
                re.search(
                    r"(\d+)",
                    x["display_name"]
                ).group(1)
            )

            if re.search(
                r"(\d+)",
                x["display_name"]
            )

            else 999

        )

    return templates.TemplateResponse(
        request=request,
        name="bom/production_plan.html",
        context={
            "grouped_rows": grouped_rows,
            "date_columns": date_columns,
            "year_columns": year_columns,
            "month_columns": month_columns,
            "week_columns": week_columns,
            "selected_year": selected_year,
            "selected_month": selected_month,
            "selected_week": selected_week,
            "selected_iso_year": selected_iso_year,
            "display_start": display_start,
            "display_end": display_end,
            "is_search_mode": is_search_mode,
            "keyword": keyword,
            "start_date": start_date,
            "end_date": end_date,
            "row_count": len(rows),
            "today_week": (
                today.isocalendar()[1]
            )
        }
    )

@router.post("/mrp/production-plan/upload")
async def upload_production_plan(
    request: Request,
    file: UploadFile = File(...),
    db: Session = Depends(get_db)
):

    try:

        df = pd.read_excel(
            file.file,
            header=None
        )

        header_row = 5

        product_col = None

        for col in range(len(df.columns)):

            value = str(
                df.iloc[header_row, col]
            ).strip()

            if value == "사이즈":

                product_col = col
                break

        if product_col is None:

            raise Exception(
                "사이즈 컬럼을 찾을 수 없습니다."
            )

        date_columns = {}

        for col in range(len(df.columns)):

            value = df.iloc[header_row, col]

            try:

                parsed_date = (
                    pd.to_datetime(value)
                    .date()
                )

                date_columns[col] = parsed_date

            except:
                continue

        existing_plans = {

            (
                p.plan_date,
                p.product_name
            ): p

            for p in db.query(
                ProductionPlan
            ).all()

        }

        for row_idx in range(
            header_row + 1,
            len(df)
        ):

            product = df.iloc[
                row_idx,
                product_col
            ]

            if pd.isna(product):
                continue

            product_name = (
                str(product)
                .strip()
            )

            for col, plan_date in date_columns.items():

                qty = df.iloc[
                    row_idx,
                    col
                ]

                if pd.isna(qty):
                    continue

                try:

                    qty = int(qty)

                except:
                    continue

                if qty <= 0:
                    continue

                existing = existing_plans.get(
                    (
                        plan_date,
                        product_name
                    )
                )

                if existing:

                    existing.plan_qty = qty

                else:

                    new_plan = ProductionPlan(
                        plan_date=plan_date,
                        product_name=product_name,
                        plan_qty=qty
                    )

                    db.add(new_plan)

                    existing_plans[
                        (
                            plan_date,
                            product_name
                        )
                    ] = new_plan

        db.commit()

        save_log(
            user=current_user(request),
            product=MRP_LOG_PRODUCT,
            action="MRP_PRODUCTION_PLAN_UPLOAD",
            detail="생산계획 업로드"
        )

        year = request.query_params.get("year")
        month = request.query_params.get("month")
        week = request.query_params.get("week")

        redirect_url = "/mrp/production-plan"

        params = []

        if year:
            params.append(f"year={year}")

        if month:
            params.append(f"month={month}")

        if week:
            params.append(f"week={week}")

        if params:
            redirect_url += "?" + "&".join(params)

        return RedirectResponse(
            url=redirect_url,
            status_code=303
        )

    except Exception as e:

        return templates.TemplateResponse(
            request=request,
            name="bom/production_plan.html",
            context={
                "rows": [],
                "date_columns": [],
                "row_count": 0,
                "error": str(e)
            }
        )
    
@router.get("/mrp/production-plan/download")
def download_production_plan(
    request: Request,
    year: int,
    month: int,
    db: Session = Depends(get_db)
):

    wb = Workbook()
    ws = wb.active

    ws.title = "생산계획"

    products = (
        db.query(BOM.product_name)
        .distinct()
        .order_by(BOM.product_name)
        .all()
    )

    plans = (
        db.query(ProductionPlan)
        .all()
    )

    plan_map = {

        (
            p.product_name,
            p.plan_date
        ): p.plan_qty

        for p in plans

    }

    selected_year = year or datetime.now().year
    selected_month = month or datetime.now().month

    last_day = calendar.monthrange(
        selected_year,
        selected_month
    )[1]

    date_list = [

        date(
            year,
            month,
            day
        )

        for day in range(
            1,
            last_day + 1
        )

    ]

    dark_fill = PatternFill(
        "solid",
        fgColor="002060"
    )

    red_fill = PatternFill(
        "solid",
        fgColor="FF0000"
    )

    white_font = Font(
        color="FFFFFF",
        bold=True
    )

    center = Alignment(
        horizontal="center",
        vertical="center"
    )

    thin = Border(

        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")

    )

    ws["B3"] = (
        f"생산 계획_{year}.{month:02d}"
    )

    ws["B3"].font = Font(
        bold=True,
        size=14
    )

    header_row = 6

    headers = [
        "MODEL",
        "사이즈",
        "출고합계",
        "생산 목표"
    ]

    for col_idx, text in enumerate(
        headers,
        start=2
    ):

        cell = ws.cell(
            header_row,
            col_idx
        )

        cell.value = text
        cell.fill = dark_fill
        cell.font = white_font
        cell.alignment = center
        cell.border = thin

    start_col = 6

    current_week = None
    week_start_col = None

    for idx, target_date in enumerate(
        date_list
    ):

        col = start_col + idx

        week_no = (
            target_date
            .isocalendar()
            .week
        )

        if current_week is None:

            current_week = week_no
            week_start_col = col

        elif current_week != week_no:

            ws.merge_cells(
                start_row=5,
                start_column=week_start_col,
                end_row=5,
                end_column=col - 1
            )

            ws.cell(
                5,
                week_start_col
            ).value = f"W{current_week}"

            current_week = week_no
            week_start_col = col

        date_cell = ws.cell(
            header_row,
            col
        )

        date_cell.value = target_date
        date_cell.number_format = "yyyy-mm-dd"

        if target_date.weekday() >= 5:

            date_cell.fill = red_fill

        else:

            date_cell.fill = dark_fill

        date_cell.font = white_font
        date_cell.alignment = center
        date_cell.border = thin

    if week_start_col:

        ws.merge_cells(
            start_row=5,
            start_column=week_start_col,
            end_row=5,
            end_column=start_col + len(date_list) - 1
        )

        ws.cell(
            5,
            week_start_col
        ).value = f"W{current_week}"

    grouped = {}

    for product in products:

        product_name = product[0]

        match = re.match(
            r"(.+?)_([0-9]+호)$",
            product_name
        )

        if match:

            group_name = (
                match.group(1)
            )

        else:

            group_name = (
                product_name
            )

        grouped.setdefault(
            group_name,
            []
        ).append(
            product_name
        )

    row_idx = 7

    for group_name, items in grouped.items():

        start_row = row_idx

        for product_name in items:

            ws.cell(
                row_idx,
                3
            ).value = product_name

            ws.cell(
                row_idx,
                3
            ).border = thin

            for idx, target_date in enumerate(
                date_list
            ):

                qty = plan_map.get(
                    (
                        product_name,
                        target_date
                    ),
                    ""
                )

                cell = ws.cell(
                    row_idx,
                    start_col + idx
                )

                cell.value = qty
                cell.border = thin
                cell.alignment = center

            row_idx += 1

        end_row = row_idx - 1

        ws.merge_cells(
            start_row=start_row,
            start_column=2,
            end_row=end_row,
            end_column=2
        )

        model_cell = ws.cell(
            start_row,
            2
        )

        model_cell.value = group_name
        model_cell.alignment = center
        model_cell.border = thin

    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 12

    for col in range(
        start_col,
        start_col + len(date_list)
    ):

        ws.column_dimensions[
            get_column_letter(col)
        ].width = 11

    output = io.BytesIO()

    wb.save(output)

    output.seek(0)

    save_log(
        user=current_user(request),
        product=MRP_LOG_PRODUCT,
        action="MRP_PRODUCTION_PLAN_DOWNLOAD",
        detail=f"생산계획 다운로드: {year}-{month:02d}"
    )

    return StreamingResponse(

        output,

        media_type=
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",

        headers={

            "Content-Disposition":
            f"attachment; filename=production_plan_{year}_{month:02d}.xlsx"

        }

    )

@router.get("/mrp/result")
def mrp_result(
    request: Request,
    db: Session = Depends(get_db),
):

    year = request.query_params.get("year")
    month = request.query_params.get("month")
    week = request.query_params.get("week")

    year = int(year) if year else None
    month = int(month) if month else None
    week = int(week) if week else None

    q = request.query_params.get(
        "q",
        ""
    )

    shortage_only = (
        request.query_params.get(
            "shortage_only"
        )
        == "true"
    )

    plans, rows, shortage_count, _ = calculate_mrp(db)

    filtered_rows = filter_mrp_rows(
        rows,
        q=q,
        shortage_only=shortage_only,
    )
    inventory_count = db.query(Inventory).count()
    summary = build_mrp_summary(
        plans,
        rows,
        filtered_rows,
        inventory_count,
    )

    plans_for_filter = (
        db.query(ProductionPlan)
        .all()
    )

    years = sorted(
        {
            p.plan_date.year
            for p in plans_for_filter
        },
        reverse=True
    )

    weeks = sorted(
        {
            (
                p.plan_date.year,
                p.plan_date.isocalendar().week
            )
            for p in plans_for_filter
        },
        reverse=True
    )

    months = sorted(
        {
            p.plan_date.month
            for p in plans_for_filter
        }
    )

    return templates.TemplateResponse(
        request=request,
        name="bom/mrp_result.html",
        context={
            "plans": plans,
            "rows": filtered_rows,
            "years": years,
            "weeks": weeks,
            "months": months,
            "selected_year": year,
            "selected_month": month,
            "selected_week": week,
            "q": q,
            "shortage_only": shortage_only,
            "shortage_count": shortage_count,
            "summary": summary,
        },
    )

@router.get("/mrp/result/download")
def download_mrp_result(
    request: Request,
    db: Session = Depends(get_db),
):

    year = request.query_params.get("year")
    month = request.query_params.get("month")
    week = request.query_params.get("week")

    year = int(year) if year else None
    month = int(month) if month else None
    week = int(week) if week else None

    q = request.query_params.get(
        "q",
        ""
    )

    shortage_only = (
        request.query_params.get(
            "shortage_only"
        )
        == "true"
    )

    _, rows, _, _ = calculate_mrp(db)

    rows = filter_mrp_rows(
        rows,
        q=q,
        shortage_only=shortage_only,
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "MRP Result"

    headers = [
        "우선순위",
        "품목코드",
        "품목명",
        "적용모델",
        "소요수량",
        "재고수량",
        "부족수량",
        "업체명",
        "LT(주)",
        "MOQ",
        "부족발생일",
        "발주필요일",
        "충족율(%)",
        "최소발주수량"
    ]

    header_fill = PatternFill(
        "solid",
        fgColor="2C3E50"
    )

    header_font = Font(
        color="FFFFFF",
        bold=True
    )

    center = Alignment(
        horizontal="center",
        vertical="center"
    )

    thin_border = Border(
        left=Side(style="thin", color="D9E2EC"),
        right=Side(style="thin", color="D9E2EC"),
        top=Side(style="thin", color="D9E2EC"),
        bottom=Side(style="thin", color="D9E2EC"),
    )

    for col_num, header in enumerate(
        headers,
        start=1
    ):

        cell = ws.cell(
            row=1,
            column=col_num
        )

        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = thin_border

    for row_idx, row in enumerate(
        rows,
        start=2
    ):

        ws.cell(
            row=row_idx,
            column=1
        ).value = row["priority"]

        item_code_cell = ws.cell(
            row=row_idx,
            column=2
        )

        item_code_cell.value = row["item_code"]

        ws.cell(
            row=row_idx,
            column=3
        ).value = row["item_name"]

        ws.cell(
            row=row_idx,
            column=4
        ).value = ", ".join(
            row["product_names"]
        )

        ws.cell(
            row=row_idx,
            column=5
        ).value = row["required_qty"]

        ws.cell(
            row=row_idx,
            column=6
        ).value = row["stock_qty"]

        shortage_cell = ws.cell(
            row=row_idx,
            column=7
        )

        shortage_cell.value = row["shortage_qty"]

        if row["shortage_qty"] > 0:

            shortage_cell.font = Font(
                color="C00000",
                bold=True
            )

        ws.cell(
            row=row_idx,
            column=8
        ).value = row["supplier"]

        ws.cell(
            row=row_idx,
            column=9
        ).value = row["lead_time_week"]

        ws.cell(
            row=row_idx,
            column=10
        ).value = row["moq"]

        need_date_cell = ws.cell(
            row=row_idx,
            column=11
        )

        if row["need_date"]:

            need_date_cell.value = row["need_date"]
            need_date_cell.number_format = "yyyy-mm-dd"

        order_date_cell = ws.cell(
            row=row_idx,
            column=12
        )

        if row["order_date"]:

            order_date_cell.value = row["order_date"]
            order_date_cell.number_format = "yyyy-mm-dd"

        coverage_cell = ws.cell(
            row=row_idx,
            column=13
        )

        coverage_cell.value = round(
            row["coverage_rate"],
            1
        )

        ws.cell(
            row=row_idx,
            column=14
        ).value = row[
            "recommended_order_qty"
        ]

        note = row.get(
            "note",
            ""
        )

        if note:

            item_code_cell.comment = Comment(
                note,
                "MRP"
            )

        for col in range(
            1,
            15
        ):

            cell = ws.cell(
                row=row_idx,
                column=col
            )

            cell.alignment = center
            cell.border = thin_border

    ws.freeze_panes = "A2"

    ws.auto_filter.ref = (
        f"A1:N{ws.max_row}"
    )

    for column in ws.columns:

        max_length = 0

        column_letter = (
            get_column_letter(
                column[0].column
            )
        )

        for cell in column:

            try:

                if cell.value:

                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            except:
                pass

        ws.column_dimensions[
            column_letter
        ].width = min(
            max_length + 5,
            50
        )

    output = BytesIO()

    wb.save(output)

    output.seek(0)

    save_log(
        user=current_user(request),
        product=MRP_LOG_PRODUCT,
        action="MRP_RESULT_DOWNLOAD",
        detail=f"MRP Result 다운로드: {len(rows)}건"
    )

    return StreamingResponse(
        output,
        media_type=
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition":
            "attachment; filename=mrp_result.xlsx"
        }
    )

@router.get("/mrp/material-master")
def material_master_page(
    request: Request,
    db: Session = Depends(get_db),
):

    rows = (
        db.query(MaterialMaster)
        .order_by(
            MaterialMaster.item_code
        )
        .all()
    )

    return templates.TemplateResponse(
        request=request,
        name="bom/material_master.html",
        context={
            "rows": rows
        }
    )

@router.post("/mrp/material-master/add")
def add_material_master(
    request: Request,

    item_code: str = Form(...),
    item_name: str = Form(...),
    supplier: str = Form(""),
    lead_time_week: int = Form(0),
    moq: int = Form(0),

    db: Session = Depends(get_db)

):

    existing = (
        db.query(MaterialMaster)
        .filter(
            MaterialMaster.item_code
            == item_code
        )
        .first()
    )

    if existing:

        existing.item_name = item_name
        existing.supplier = supplier
        existing.lead_time_week = lead_time_week
        existing.moq = moq

    else:

        db.add(

            MaterialMaster(
                item_code=item_code,
                item_name=item_name,
                supplier=supplier,
                lead_time_week=lead_time_week,
                moq=moq
            )

        )

    db.commit()

    save_log(
        user=current_user(request),
        product=MRP_LOG_PRODUCT,
        action="MRP_MATERIAL_MASTER_ADD",
        serial=item_code,
        detail=f"자재 기준정보 저장: {item_name} / 업체 {supplier or '-'} / LT {lead_time_week}주 / MOQ {moq}"
    )

    return RedirectResponse(
        "/mrp/material-master",
        status_code=303
    )

@router.post(
    "/mrp/material-master/update/{row_id}"
)
async def update_material_master(

    row_id: int,
    request: Request,
    db: Session = Depends(get_db)

):

    row = db.get(
        MaterialMaster,
        row_id
    )

    if not row:
        return {"success": False}

    data = await request.form()

    field = data.get("field")
    value = data.get("value")

    if field == "supplier":

        row.supplier = value

    elif field == "lead_time_week":

        row.lead_time_week = (
            int(value)
            if str(value).strip()
            else 0
        )

    elif field == "moq":

        value = str(value).strip().upper()

        if value in ["", "-"]:

            row.moq = 0

        elif value.endswith("K"):

            row.moq = int(
                float(
                    value.replace("K", "")
                ) * 1000
            )

        else:

            row.moq = int(
                float(value)
            )

    db.commit()

    save_log(
        user=current_user(request),
        product=MRP_LOG_PRODUCT,
        action="MRP_MATERIAL_MASTER_UPDATE",
        serial=row.item_code,
        detail=f"자재 기준정보 수정: {row.item_name} / {field}={value}"
    )

    return {"success": True}

@router.get("/mrp/material-master/download")
def download_material_master(
    request: Request,
    db: Session = Depends(get_db)
):

    rows = (
        db.query(MaterialMaster)
        .order_by(
            MaterialMaster.item_code
        )
        .all()
    )

    data = []

    for row in rows:

        data.append(
            {
                "품목코드": row.item_code,
                "품목명": row.item_name,
                "업체명": row.supplier,
                "LT(주)": row.lead_time_week,
                "MOQ": row.moq,
            }
        )

    output = BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl"
    ) as writer:

        pd.DataFrame(data).to_excel(
            writer,
            index=False,
            sheet_name="Material Master"
        )

    output.seek(0)

    save_log(
        user=current_user(request),
        product=MRP_LOG_PRODUCT,
        action="MRP_MATERIAL_MASTER_DOWNLOAD",
        detail=f"자재 기준정보 다운로드: {len(rows)}건"
    )

    return StreamingResponse(
        output,
        media_type=
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition":
            "attachment; filename=material_master.xlsx"
        }
    )

@router.post("/mrp/material-master/upload")
async def upload_material_master(
    request: Request,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):

    try:

        df = pd.read_excel(file.file)

        leadtime_col = None

        for col in df.columns:

            col_name = str(col).strip()

            if col_name in [
                "LT(주)",
                "LT",
                "L/T",
                "L/Time",
                "리드타임",
                "Lead Time"
            ]:
                leadtime_col = col_name
                break

        if leadtime_col is None:

            raise Exception(
                "리드타임 컬럼이 없습니다."
            )
        
        required_columns = [
            "품목코드",
            "품목명",
            "업체명",
            "MOQ"
        ]

        for col in required_columns:

            if col not in df.columns:

                raise Exception(
                    f"{col} 컬럼이 없습니다."
                )

        for _, row in df.iterrows():

            if pd.isna(row["품목코드"]):
                continue

            item_code = str(
                row["품목코드"]
            ).strip()

            if item_code == "":
                continue

            existing = (
                db.query(MaterialMaster)
                .filter(
                    MaterialMaster.item_code
                    == item_code
                )
                .first()
            )

            if existing:

                existing.item_name = str(
                    row["품목명"]
                ).strip()

                existing.supplier = str(
                    row["업체명"]
                ).strip()

                existing.lead_time_week = int(
                    row[leadtime_col]
                    if pd.notna(row[leadtime_col])
                    else 0
                )

                existing.moq = parse_moq(
                    row["MOQ"]
                )

            else:

                db.add(

                    MaterialMaster(

                        item_code=item_code,

                        item_name=str(
                            row["품목명"]
                        ).strip(),

                        supplier=str(
                            row["업체명"]
                        ).strip(),

                        lead_time_week=int(
                            row[leadtime_col]
                            if pd.notna(row[leadtime_col])
                            else 0
                        ),

                        moq=parse_moq(
                            row["MOQ"]
                        )

                    )

                )

        db.commit()

        save_log(
            user=current_user(request),
            product=MRP_LOG_PRODUCT,
            action="MRP_MATERIAL_MASTER_UPLOAD",
            detail="자재 기준정보 업로드"
        )

        return RedirectResponse(
            "/mrp/material-master",
            status_code=303
        )

    except Exception as e:

        return templates.TemplateResponse(
            request=request,
            name="bom/material_master.html",
            context={
                "rows": (
                    db.query(MaterialMaster)
                    .all()
                ),
                "error": str(e)
            },
            status_code=400
        )
    
@router.get(
    "/mrp/material-note/{item_code}"
)
def get_material_note(
    item_code: str,
    db: Session = Depends(get_db)
):

    row = (
        db.query(MaterialNote)
        .filter(
            MaterialNote.item_code
            == item_code
        )
        .first()
    )

    return {
        "note":
            row.note
            if row
            else ""
    }

@router.post("/mrp/material-note/save")
async def save_material_note(
    request: Request,
    db: Session = Depends(get_db)
):

    data = await request.json()

    item_code = data.get(
        "item_code"
    )

    note = data.get(
        "note",
        ""
    )

    row = (
        db.query(MaterialNote)
        .filter(
            MaterialNote.item_code
            == item_code
        )
        .first()
    )

    if row:

        row.note = note

    else:

        db.add(
            MaterialNote(
                item_code=item_code,
                note=note
            )
        )

    db.commit()

    save_log(
        user=current_user(request),
        product=MRP_LOG_PRODUCT,
        action="MRP_MATERIAL_NOTE_SAVE",
        serial=item_code,
        detail="자재 노트 저장"
    )

    return {
        "success": True
    }
