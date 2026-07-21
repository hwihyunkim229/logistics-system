from datetime import date
from io import BytesIO
from math import ceil
import pandas as pd
from openpyxl.worksheet.datavalidation import DataValidation
from fastapi import (
    APIRouter,
    Body,
    Depends,
    File,
    Request,
    UploadFile,
)
from fastapi.responses import (
    JSONResponse,
    StreamingResponse,
)
from fastapi.templating import Jinja2Templates
from sqlalchemy import or_
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from app.database import SessionLocal
from app.models.item_master import ItemMaster
from app.models.rental_category import RentalCategory
from app.models.rental_movement import RentalMovement
from app.models.rental_stock import RentalStock
from app.utils.downloads import download_content_disposition
from app.utils.logger import save_log

router = APIRouter(prefix="/rental")

templates = Jinja2Templates(
    directory="app/templates"
)


HQ = {
    "경영지원본부": [
        "구매/자재/물류팀",
        "디자인팀",
        "인사총무팀",
    ],
    "경영기획본부": [
        "회계팀",
        "IR/자금팀",
    ],
    "대외협력본부": [
        "MA팀",
        "PR팀",
    ],
    "제품개발본부": [
        "제품기획/PM팀",
        "서비스운영팀",
        "서비스기획팀",
        "인허가팀",
        "정보보안팀",
    ],
    "연구개발본부": [
        "연구기획팀",
        "VSP팀",
        "Server팀",
        "Mobile SW팀",
        "HW/FW팀",
        "기구팀",
    ],
    "품질본부": [
        "QA팀",
        "QC팀",
        "SW품질팀",
    ],
    "생산본부": [
        "생산관리팀",
        "생산팀",
    ],
    "사업본부": [
        "사업개발팀",
        "전략마케팅팀",
        "마케팅팀",
    ],
    "과학본부": [
        "임상연구1팀",
        "임상연구2팀",
    ],
}

SIZES = [
    f"{number}호"
    for number in range(7, 14)
]

def get_db():
    db = SessionLocal()

    try:
        yield db
    finally:
        db.close()

def user(request):
    return (
        request.session.get("user")
        or "system"
    )

def xlsx(
    rows,
    filename,
    sheet,
    columns=None,
    configure=None,
    download_label=None,
):
    output = BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl",
    ) as writer:
        pd.DataFrame(rows, columns=columns).to_excel(
            writer,
            index=False,
            sheet_name=sheet,
        )

        if configure:
            configure(
                writer.book,
                writer.sheets[sheet],
            )

    output.seek(0)

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                download_content_disposition(
                    download_label or sheet,
                    filename.rsplit(".", 1)[0],
                )
            )
        },
    )


def configure_rental_history_excel(
    workbook,
    worksheet,
    categories,
):
    list_sheet = workbook.create_sheet("선택목록")

    list_sheet.cell(1, 1, "구분")
    list_sheet.cell(2, 1, "출고")
    list_sheet.cell(3, 1, "입고")

    list_sheet.cell(1, 2, "분류")
    for row_index, category in enumerate(categories, start=2):
        list_sheet.cell(row_index, 2, category)

    list_sheet.cell(1, 3, "사이즈")
    for row_index, size in enumerate(SIZES, start=2):
        list_sheet.cell(row_index, 3, size)

    hq_start_column = 4
    for column_index, (headquarters, departments) in enumerate(
        HQ.items(),
        start=hq_start_column,
    ):
        list_sheet.cell(1, column_index, headquarters)
        for row_index, department in enumerate(departments, start=2):
            list_sheet.cell(row_index, column_index, department)

    hq_end_column = hq_start_column + len(HQ) - 1
    category_end_row = max(2, len(categories) + 1)
    size_end_row = len(SIZES) + 1
    hq_start_letter = list_sheet.cell(1, hq_start_column).column_letter
    hq_end_letter = list_sheet.cell(1, hq_end_column).column_letter

    def add_validation(column, formula, prompt):
        validation = DataValidation(
            type="list",
            formula1=formula,
            allow_blank=False,
        )
        validation.error = "목록에서 값을 선택해 주세요."
        validation.errorTitle = "입력값 확인"
        validation.prompt = prompt
        validation.promptTitle = "선택 입력"
        validation.showErrorMessage = True
        validation.showInputMessage = True
        worksheet.add_data_validation(validation)
        validation.add(f"{column}2:{column}1001")

    add_validation("B", "='선택목록'!$A$2:$A$3", "입고 또는 출고를 선택하세요.")
    add_validation(
        "C",
        f"='선택목록'!$B$2:$B${category_end_row}",
        "재고 현황에 등록된 분류를 선택하세요.",
    )
    add_validation(
        "D",
        (
            f"='선택목록'!${hq_start_letter}$1:"
            f"${hq_end_letter}$1"
        ),
        "본부를 선택하세요.",
    )

    department_formula = (
        f"=OFFSET('선택목록'!${hq_start_letter}$1,1,"
        f"MATCH($D2,'선택목록'!${hq_start_letter}$1:"
        f"${hq_end_letter}$1,0)-1,"
        f"COUNTA(OFFSET('선택목록'!${hq_start_letter}$1,0,"
        f"MATCH($D2,'선택목록'!${hq_start_letter}$1:"
        f"${hq_end_letter}$1,0)-1,100,1))-1,1)"
    )
    add_validation(
        "E",
        department_formula,
        "본부를 먼저 선택한 뒤 해당 부서를 선택하세요.",
    )
    add_validation(
        "J",
        f"='선택목록'!$C$2:$C${size_end_row}",
        "사이즈를 선택하세요.",
    )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:M{max(1, worksheet.max_row)}"
    widths = {
        "A": 11,
        "B": 10,
        "C": 18,
        "D": 18,
        "E": 22,
        "F": 18,
        "G": 28,
        "H": 10,
        "I": 14,
        "J": 10,
        "K": 20,
        "L": 14,
        "M": 28,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width

    list_sheet.sheet_state = "hidden"

def category_names(
    db: Session,
):
    rows = (
        db.query(RentalCategory)
        .order_by(RentalCategory.id)
        .all()
    )

    return [
        row.name
        for row in rows
    ]

@router.get("")
def index(
    request: Request,
    keyword: str = "",
    category: str = "",
    page: int = 1,
    db: Session = Depends(get_db),
):
    categories = category_names(db)

    if category not in categories:
        category = (
            categories[0]
            if categories
            else ""
        )

    if category:
        query = (
            db.query(RentalStock)
            .filter(
                RentalStock.category
                == category
            )
        )
    else:
        query = (
            db.query(RentalStock)
            .filter(False)
        )

    if keyword:
        query = query.filter(
            or_(
                RentalStock.item_code.contains(
                    keyword
                ),
                RentalStock.item_name.contains(
                    keyword
                ),
            )
        )

    total = query.count()

    items = (
        query
        .order_by(RentalStock.item_code)
        .offset((page - 1) * 50)
        .limit(50)
        .all()
    )

    item_master_rows = (
        db.query(ItemMaster)
        .filter(
            ItemMaster.item_code.isnot(None)
        )
        .order_by(ItemMaster.item_code)
        .all()
    )

    item_master_list = [
        {
            "code": row.item_code,
            "name": row.item_name,
        }
        for row in item_master_rows
    ]

    return templates.TemplateResponse(
        request=request,
        name="rental/index.html",
        context={
            "items": items,
            "keyword": keyword,
            "category": category,
            "categories": categories,
            "item_master_list": item_master_list,
            "page": page,
            "total_pages": max(
                1,
                ceil(total / 50),
            ),
            "headquarters": HQ,
            "sizes": SIZES,
            "today": date.today().isoformat(),
        },
    )

@router.post("/categories")
def add_category(
    data: dict = Body(...),
    db: Session = Depends(get_db),
):
    name = " ".join(
        str(
            data.get("name", "")
        ).split()
    )

    if not name:
        return JSONResponse(
            {
                "status": "error",
                "message": (
                    "분류명을 입력해 주세요."
                ),
            },
            status_code=400,
        )

    if len(name) > 50:
        return JSONResponse(
            {
                "status": "error",
                "message": (
                    "분류명은 50자 이내로 "
                    "입력해 주세요."
                ),
            },
            status_code=400,
        )

    existing_category = (
        db.query(RentalCategory)
        .filter(
            RentalCategory.name == name
        )
        .first()
    )

    if existing_category:
        return JSONResponse(
            {
                "status": "error",
                "message": (
                    "이미 등록된 분류입니다."
                ),
            },
            status_code=409,
        )

    db.add(
        RentalCategory(
            name=name
        )
    )

    db.commit()

    return {
        "status": "success",
        "name": name,
    }

@router.post("/items")
def add_item(
    request: Request,
    data: dict = Body(...),
    db: Session = Depends(get_db),
):
    category = str(
        data.get("category", "")
    ).strip()

    item_code = str(
        data.get("item_code", "")
    ).strip()

    item_name = str(
        data.get("item_name", "")
    ).strip()

    try:
        qty = int(
            data.get("qty", 0)
        )
    except (TypeError, ValueError):
        return JSONResponse(
            {
                "status": "error",
                "message": (
                    "수량은 숫자로 입력해 주세요."
                ),
            },
            status_code=400,
        )

    if category not in category_names(db):
        return JSONResponse(
            {
                "status": "error",
                "message": (
                    "대여 분류를 확인해 주세요."
                ),
            },
            status_code=400,
        )

    master = (
        db.query(ItemMaster)
        .filter(
            ItemMaster.item_code
            == item_code
        )
        .first()
    )

    if (
        master
        and master.item_name
    ):
        item_name = master.item_name

    if (
        not item_code
        or not item_name
        or qty < 0
    ):
        return JSONResponse(
            {
                "status": "error",
                "message": (
                    "품목코드, 품명, 수량을 "
                    "올바르게 입력해 주세요."
                ),
            },
            status_code=400,
        )

    existing_item = (
        db.query(RentalStock)
        .filter(
            RentalStock.category
            == category,
            RentalStock.item_code
            == item_code,
        )
        .first()
    )

    if existing_item:
        return JSONResponse(
            {
                "status": "error",
                "message": (
                    "해당 분류에 이미 등록된 "
                    "품목코드입니다."
                ),
            },
            status_code=409,
        )

    db.add(
        RentalStock(
            category=category,
            item_code=item_code,
            item_name=item_name,
            qty=qty,
        )
    )

    try:
        db.commit()
    except IntegrityError:
        db.rollback()

        return JSONResponse(
            {
                "status": "error",
                "message": (
                    "이미 등록된 품목코드입니다."
                ),
            },
            status_code=409,
        )

    save_log(
        user=user(request),
        product="대여 관리",
        action="RENTAL_ITEM_CREATE",
        serial=item_code,
        detail=(
            f"{category} 신규 품목 등록: "
            f"{item_name} / {qty} EA"
        ),
    )

    return {
        "status": "success"
    }


@router.patch(
    "/items/{item_id}/qty"
)
def update_item_qty(
    item_id: int,
    request: Request,
    data: dict = Body(...),
    db: Session = Depends(get_db),
):
    raw_qty = str(
        data.get("qty", "")
    ).strip()

    if not raw_qty.isdigit():
        return JSONResponse(
            {
                "status": "error",
                "message": (
                    "수량은 0 이상의 정수로 "
                    "입력해 주세요."
                ),
            },
            status_code=400,
        )

    qty = int(raw_qty)

    item = db.get(
        RentalStock,
        item_id,
    )

    if not item:
        return JSONResponse(
            {
                "status": "error",
                "message": (
                    "대여 재고를 찾을 수 없습니다."
                ),
            },
            status_code=404,
        )

    old_qty = item.qty

    item.qty = qty

    db.commit()

    save_log(
        user=user(request),
        product="대여 관리",
        action="RENTAL_QTY_UPDATE",
        serial=item.item_code,
        detail=(
            f"{item.category} 수량 변경: "
            f"{old_qty} → {qty} EA"
        ),
    )

    return {
        "status": "success",
        "qty": qty,
    }


@router.post("/upload-excel")
async def upload(
    request: Request,
    category: str = "",
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    categories = category_names(db)

    if category not in categories:
        return JSONResponse(
            {
                "status": "error",
                "message": (
                    "대여 분류를 확인해 주세요."
                ),
            },
            status_code=400,
        )

    try:
        contents = await file.read()

        frame = pd.read_excel(
            BytesIO(contents)
        ).fillna("")
    except Exception:
        return JSONResponse(
            {
                "status": "error",
                "message": (
                    "엑셀 파일을 읽을 수 없습니다."
                ),
            },
            status_code=400,
        )

    columns = {
        str(column)
        .strip()
        .replace(" ", ""): column
        for column in frame.columns
    }

    required_columns = (
        "품목코드",
        "품명",
        "수량",
    )

    if not all(
        column in columns
        for column in required_columns
    ):
        return JSONResponse(
            {
                "status": "error",
                "message": (
                    "품목코드, 품명, 수량 "
                    "컬럼이 필요합니다."
                ),
            },
            status_code=400,
        )

    count = 0

    try:
        for _, excel_row in frame.iterrows():
            code = str(
                excel_row[
                    columns["품목코드"]
                ]
            ).strip()

            name = str(
                excel_row[
                    columns["품명"]
                ]
            ).strip()

            if (
                not code
                or not name
            ):
                continue

            qty = int(
                float(
                    excel_row[
                        columns["수량"]
                    ]
                    or 0
                )
            )

            row = (
                db.query(RentalStock)
                .filter_by(
                    item_code=code,
                    category=category,
                )
                .first()
            )

            if row:
                row.item_name = name
                row.qty = qty
            else:
                db.add(
                    RentalStock(
                        item_code=code,
                        item_name=name,
                        category=category,
                        qty=qty,
                    )
                )

            count += 1

        db.commit()

    except (ValueError, TypeError):
        db.rollback()

        return JSONResponse(
            {
                "status": "error",
                "message": (
                    "수량은 숫자로 입력해 주세요."
                ),
            },
            status_code=400,
        )

    save_log(
        user=user(request),
        product="대여 관리",
        action="RENTAL_UPLOAD_EXCEL",
        detail=(
            f"대여 재고 업로드: "
            f"{count}건"
        ),
    )

    return {
        "status": "success",
        "count": count,
    }


@router.get("/download-excel")
def stock_excel(
    request: Request,
    category: str = "",
    db: Session = Depends(get_db),
):
    categories = category_names(db)

    if category not in categories:
        if not categories:
            return JSONResponse(
                {
                    "status": "error",
                    "message": (
                        "먼저 대여 분류를 "
                        "추가해 주세요."
                    ),
                },
                status_code=400,
            )

        category = categories[0]

    rows = (
        db.query(RentalStock)
        .filter(
            RentalStock.category
            == category
        )
        .order_by(RentalStock.item_code)
        .all()
    )

    save_log(
        user=user(request),
        product="대여 관리",
        action="RENTAL_DOWNLOAD_EXCEL",
        detail=(
            f"{category} 대여 재고 다운로드: "
            f"{len(rows)}건"
        ),
    )

    excel_rows = [
        {
            "품목코드": row.item_code,
            "품명": row.item_name,
            "수량": row.qty,
        }
        for row in rows
    ]

    return xlsx(
        excel_rows,
        "rental_stock.xlsx",
        "대여 재고",
        download_label="대여 재고 현황",
    )


@router.post("/move-out")
def move_out(
    request: Request,
    data: dict = Body(...),
    db: Session = Depends(get_db),
):
    try:
        ids = list(
            {
                int(item_id)
                for item_id in data.get(
                    "ids",
                    [],
                )
            }
        )

        qty = int(
            data.get("qty", 1)
        )

        issued = date.fromisoformat(
            data.get(
                "issue_date",
                "",
            )
        )

    except (ValueError, TypeError):
        return JSONResponse(
            {
                "status": "error",
                "message": (
                    "수량과 불출 일자를 "
                    "확인해 주세요."
                ),
            },
            status_code=400,
        )

    headquarters = data.get(
        "headquarters",
        "",
    )

    department = data.get(
        "department",
        "",
    )

    requester = data.get(
        "requester",
        "",
    ).strip()

    serial_number = data.get(
        "serial_number",
        "",
    ).strip()

    size = data.get("size")

    if (
        not ids
        or qty < 1
        or headquarters not in HQ
        or department
        not in HQ[headquarters]
        or size not in SIZES
        or not requester
        or not serial_number
    ):
        return JSONResponse(
            {
                "status": "error",
                "message": (
                    "필수 항목을 모두 올바르게 "
                    "입력해 주세요."
                ),
            },
            status_code=400,
        )

    items = (
        db.query(RentalStock)
        .filter(
            RentalStock.id.in_(ids)
        )
        .all()
    )

    if len(items) != len(ids):
        return JSONResponse(
            {
                "status": "error",
                "message": (
                    "선택 품목을 찾을 수 없습니다."
                ),
            },
            status_code=404,
        )

    short_item = next(
        (
            item
            for item in items
            if item.qty < qty
        ),
        None,
    )

    if short_item:
        return JSONResponse(
            {
                "status": "error",
                "message": (
                    f"{short_item.item_code} "
                    "재고가 부족합니다."
                ),
            },
            status_code=400,
        )

    remark = data.get(
        "remark",
        "",
    ).strip()

    username = user(request)

    for item in items:
        item.qty -= qty

        db.add(
            RentalMovement(
                movement_type="OUT",
                item_code=item.item_code,
                item_name=item.item_name,
                category=item.category,
                qty=qty,
                headquarters=headquarters,
                department=department,
                requester=requester,
                size=size,
                serial_number=serial_number,
                issue_date=issued,
                remark=remark,
                user=username,
            )
        )

    db.commit()

    save_log(
        user=username,
        product="대여 관리",
        action="RENTAL_MOVE_OUT",
        serial=", ".join(
            item.item_code
            for item in items
        ),
        detail=(
            f"대여 재고 출고: "
            f"{len(items)}개 품목 × {qty}"
        ),
    )

    return {
        "status": "success"
    }


@router.get("/history")
def history(
    request: Request,
    keyword: str = "",
    page: int = 1,
    db: Session = Depends(get_db),
):
    query = db.query(
        RentalMovement
    )

    if keyword:
        query = query.filter(
            or_(
                RentalMovement.item_code.contains(
                    keyword
                ),
                RentalMovement.item_name.contains(
                    keyword
                ),
                RentalMovement.requester.contains(
                    keyword
                ),
                RentalMovement.serial_number.contains(
                    keyword
                ),
                RentalMovement.headquarters.contains(
                    keyword
                ),
                RentalMovement.department.contains(
                    keyword
                ),
            )
        )

    total = query.count()

    rows = (
        query
        .order_by(
            RentalMovement.issue_date.desc(),
            RentalMovement.id.desc(),
        )
        .offset((page - 1) * 50)
        .limit(50)
        .all()
    )

    return templates.TemplateResponse(
        request=request,
        name="rental/history.html",
        context={
            "history": rows,
            "keyword": keyword,
            "page": page,
            "total_pages": max(
                1,
                ceil(total / 50),
            ),
        },
    )


@router.post("/history/upload-excel")
async def history_upload_excel(
    request: Request,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    try:
        contents = await file.read()
        frame = pd.read_excel(BytesIO(contents)).fillna("")
    except Exception:
        return JSONResponse(
            {
                "status": "error",
                "message": "엑셀 파일을 읽을 수 없습니다.",
            },
            status_code=400,
        )

    columns = {
        str(column).strip().replace(" ", ""): column
        for column in frame.columns
    }
    required_columns = (
        "이력ID",
        "구분",
        "분류",
        "품목코드",
        "품명",
        "수량",
        "본부",
        "부서",
        "요청자",
        "사이즈",
        "S/N",
        "불출일자",
        "비고",
    )

    missing = [
        column for column in required_columns
        if column not in columns
    ]
    if missing:
        return JSONResponse(
            {
                "status": "error",
                "message": (
                    "입출고 현황에서 다운로드한 양식을 사용해 주세요. "
                    f"누락 컬럼: {', '.join(missing)}"
                ),
            },
            status_code=400,
        )

    pending = []
    skipped = 0
    errors = []

    for index, excel_row in frame.iterrows():
        excel_line = index + 2

        if str(excel_row[columns["이력ID"]]).strip():
            skipped += 1
            continue

        values = {
            key: str(excel_row[columns[key]]).strip()
            for key in required_columns
            if key not in {"이력ID", "수량", "불출일자"}
        }

        if not any(values.values()) and not str(
            excel_row[columns["수량"]]
        ).strip():
            continue

        movement_text = values["구분"].upper()
        movement_type = {
            "출고": "OUT",
            "OUT": "OUT",
            "입고": "IN",
            "IN": "IN",
        }.get(movement_text)

        try:
            qty_value = float(excel_row[columns["수량"]])
            qty = int(qty_value)
            if qty < 1 or qty != qty_value:
                raise ValueError
        except (TypeError, ValueError):
            errors.append(f"{excel_line}행: 수량은 1 이상의 정수여야 합니다.")
            continue

        try:
            raw_date = excel_row[columns["불출일자"]]
            if isinstance(raw_date, (int, float)):
                raise ValueError
            issue_date = pd.to_datetime(raw_date).date()
        except (TypeError, ValueError):
            errors.append(f"{excel_line}행: 불출일자를 확인해 주세요.")
            continue

        if movement_type is None:
            errors.append(f"{excel_line}행: 구분은 입고 또는 출고여야 합니다.")
        elif not values["분류"] or not values["품목코드"]:
            errors.append(f"{excel_line}행: 분류와 품목코드는 필수입니다.")
        elif values["본부"] not in HQ:
            errors.append(f"{excel_line}행: 본부를 확인해 주세요.")
        elif values["부서"] not in HQ[values["본부"]]:
            errors.append(f"{excel_line}행: 본부와 부서가 일치하지 않습니다.")
        elif not values["요청자"]:
            errors.append(f"{excel_line}행: 요청자는 필수입니다.")
        elif values["사이즈"] not in SIZES:
            errors.append(f"{excel_line}행: 사이즈를 확인해 주세요.")
        elif not values["S/N"]:
            errors.append(f"{excel_line}행: S/N은 필수입니다.")
        else:
            pending.append({
                "line": excel_line,
                "movement_type": movement_type,
                "category": values["분류"],
                "item_code": values["품목코드"],
                "qty": qty,
                "headquarters": values["본부"],
                "department": values["부서"],
                "requester": values["요청자"],
                "size": values["사이즈"],
                "serial_number": values["S/N"],
                "issue_date": issue_date,
                "remark": values["비고"],
            })

    if errors:
        return JSONResponse(
            {
                "status": "error",
                "message": "\n".join(errors[:10]) + (
                    f"\n외 {len(errors) - 10}건" if len(errors) > 10 else ""
                ),
            },
            status_code=400,
        )

    if not pending:
        return JSONResponse(
            {
                "status": "error",
                "message": "추가할 신규 행이 없습니다. 이력ID가 빈 행을 추가해 주세요.",
            },
            status_code=400,
        )

    keys = {
        (row["category"], row["item_code"])
        for row in pending
    }

    try:
        stocks = (
            db.query(RentalStock)
            .filter(or_(*[
                (
                    (RentalStock.category == category)
                    & (RentalStock.item_code == item_code)
                )
                for category, item_code in keys
            ]))
            .with_for_update()
            .all()
        )
        stock_map = {
            (stock.category, stock.item_code): stock
            for stock in stocks
        }

        missing_items = [
            row for row in pending
            if (row["category"], row["item_code"]) not in stock_map
        ]
        if missing_items:
            missing_text = ", ".join(
                f"{row['line']}행 {row['category']}/{row['item_code']}"
                for row in missing_items[:10]
            )
            raise ValueError(f"재고 현황에서 품목을 찾을 수 없습니다: {missing_text}")

        deltas = {key: 0 for key in keys}
        for row in pending:
            key = (row["category"], row["item_code"])
            deltas[key] += (
                row["qty"] if row["movement_type"] == "IN" else -row["qty"]
            )

        shortages = [
            f"{category}/{item_code} (현재 {stock_map[key].qty}, 변경 {delta:+d})"
            for key, delta in deltas.items()
            for category, item_code in [key]
            if stock_map[key].qty + delta < 0
        ]
        if shortages:
            raise ValueError("재고가 부족합니다: " + ", ".join(shortages[:10]))

        for key, delta in deltas.items():
            stock_map[key].qty += delta

        username = user(request)
        for row in pending:
            stock = stock_map[(row["category"], row["item_code"])]
            db.add(RentalMovement(
                movement_type=row["movement_type"],
                item_code=stock.item_code,
                item_name=stock.item_name,
                category=stock.category,
                qty=row["qty"],
                headquarters=row["headquarters"],
                department=row["department"],
                requester=row["requester"],
                size=row["size"],
                serial_number=row["serial_number"],
                issue_date=row["issue_date"],
                remark=row["remark"],
                user=username,
            ))

        db.commit()
    except ValueError as exc:
        db.rollback()
        return JSONResponse(
            {"status": "error", "message": str(exc)},
            status_code=400,
        )
    except Exception:
        db.rollback()
        return JSONResponse(
            {
                "status": "error",
                "message": "업로드 처리 중 오류가 발생했습니다. 반영된 내역은 없습니다.",
            },
            status_code=500,
        )

    save_log(
        user=user(request),
        product="대여 관리",
        action="RENTAL_HISTORY_UPLOAD_EXCEL",
        detail=f"대여 입출고 업로드: 신규 {len(pending)}건, 기존 {skipped}건 제외",
    )

    return {
        "status": "success",
        "count": len(pending),
        "skipped": skipped,
        "message": (
            f"신규 {len(pending)}건을 반영했습니다. "
            f"기존 이력 {skipped}건은 제외했습니다."
        ),
    }


@router.patch("/history/{row_id}")
def edit(
    row_id: int,
    data: dict = Body(...),
    db: Session = Depends(get_db),
):
    row = db.get(
        RentalMovement,
        row_id,
    )

    field = data.get("field")

    if (
        not row
        or field
        not in {
            "remark",
            "requester",
            "issue_date",
        }
    ):
        return JSONResponse(
            {
                "status": "error",
                "message": (
                    "수정할 수 없습니다."
                ),
            },
            status_code=400,
        )

    value = str(data.get("value", "")).strip()

    if field == "issue_date":
        try:
            row.issue_date = date.fromisoformat(value)
        except ValueError:
            return JSONResponse(
                {
                    "status": "error",
                    "message": "불출 일자를 올바른 날짜로 입력해 주세요.",
                },
                status_code=400,
            )
    else:
        setattr(row, field, value)

    db.commit()

    return {
        "status": "success"
    }


@router.get(
    "/history/download-excel"
)
def history_excel(
    request: Request,
    db: Session = Depends(get_db),
):
    categories = category_names(db)
    rows = (
        db.query(RentalMovement)
        .order_by(
            RentalMovement.issue_date.desc(),
            RentalMovement.id.desc(),
        )
        .all()
    )

    save_log(
        user=user(request),
        product="대여 관리",
        action=(
            "RENTAL_HISTORY_DOWNLOAD_EXCEL"
        ),
        detail=(
            f"대여 입출고 다운로드: "
            f"{len(rows)}건"
        ),
    )

    excel_rows = [
        {
            "이력ID": row.id,
            "구분": "입고" if row.movement_type == "IN" else "출고",
            "분류": row.category,
            "본부": row.headquarters,
            "부서": row.department,
            "품목코드": row.item_code,
            "품명": row.item_name,
            "수량": row.qty,
            "요청자": row.requester,
            "사이즈": row.size,
            "S/N": row.serial_number,
            "불출일자": (
                row.issue_date.isoformat()
            ),
            "비고": row.remark,
        }
        for row in rows
    ]

    return xlsx(
        excel_rows,
        "rental_movement_history.xlsx",
        "대여 입출고 현황",
        columns=[
            "이력ID",
            "구분",
            "분류",
            "본부",
            "부서",
            "품목코드",
            "품명",
            "수량",
            "요청자",
            "사이즈",
            "S/N",
            "불출일자",
            "비고",
        ],
        configure=lambda workbook, worksheet: (
            configure_rental_history_excel(
                workbook,
                worksheet,
                categories,
            )
        ),
        download_label="대여 입출고 현황",
    )